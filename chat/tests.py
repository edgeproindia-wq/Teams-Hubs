import io
import json

from channels.routing import URLRouter
from channels.testing import WebsocketCommunicator
from django.core.cache import cache
from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .models import Company, Group, GroupMember, Message, User
from .routing import websocket_urlpatterns


def make_company(name="Acme"):
    return Company.objects.create(name=name, slug=Company.make_slug(name))


def make_user(company, email, role=User.EMPLOYEE, first="Test", last="User", pw="Str0ng-pass-99"):
    u = User(username=email, email=email, first_name=first, last_name=last, company=company, role=role)
    u.set_password(pw)
    u.save()
    return u


class LoginTests(TestCase):
    def setUp(self):
        cache.clear()
        self.u = make_user(make_company(), "a@acme.com")

    def test_login_ok_and_logout(self):
        r = self.client.post(reverse("login"), {"username": "A@Acme.com", "password": "Str0ng-pass-99"})
        self.assertRedirects(r, reverse("dashboard"))
        self.client.post(reverse("logout"))
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 302)

    def test_wrong_password_and_lockout(self):
        for _ in range(5):
            r = self.client.post(reverse("login"), {"username": "a@acme.com", "password": "nope"})
            self.assertContains(r, "Incorrect Login ID or password")
        r = self.client.post(reverse("login"), {"username": "a@acme.com", "password": "Str0ng-pass-99"})
        self.assertContains(r, "Too many failed attempts")

    def test_deactivated_user_cannot_login(self):
        self.u.is_active = False
        self.u.save()
        r = self.client.post(reverse("login"), {"username": "a@acme.com", "password": "Str0ng-pass-99"})
        self.assertEqual(r.status_code, 200)

    def test_pages_require_login(self):
        for name in ("dashboard", "chat", "directory", "api_employees", "admin_panel"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 302 if name != "api_employees" else 401, name)

    def test_health(self):
        self.assertEqual(self.client.get("/health/").content, b"ok")


class IsolationAndApiTests(TestCase):
    def setUp(self):
        cache.clear()
        self.a, self.b = make_company("A Co"), make_company("B Co")
        self.a1 = make_user(self.a, "a1@a.com", first="Anil")
        self.a2 = make_user(self.a, "a2@a.com", first="Bina")
        self.b1 = make_user(self.b, "b1@b.com", first="Chen")
        Message.objects.create(company=self.a, sender=self.a1, recipient=self.a2, body="hello <b>bina</b>")
        Message.objects.create(company=self.b, sender=self.b1, recipient=self.b1, body="secret")
        self.client.force_login(self.a1)

    def test_directory_lists_only_own_company(self):
        names = [e["name"] for e in self.client.get(reverse("api_employees")).json()["employees"]]
        self.assertEqual(names, ["Bina User"])

    def test_search_by_name(self):
        self.assertEqual(len(self.client.get(reverse("api_employees"), {"q": "bin"}).json()["employees"]), 1)
        self.assertEqual(len(self.client.get(reverse("api_employees"), {"q": "chen"}).json()["employees"]), 0)

    def test_cannot_read_other_company_history(self):
        self.assertEqual(self.client.get(reverse("api_messages", args=[self.b1.id])).status_code, 404)

    def test_history_and_mark_read(self):
        self.client.force_login(self.a2)
        j = self.client.get(reverse("api_messages", args=[self.a1.id])).json()
        self.assertEqual(j["messages"][0]["body"], "hello <b>bina</b>")  # raw text; browser uses textContent
        self.assertIsNotNone(Message.objects.get(sender=self.a1).read_at)

    def test_online_filter(self):
        from django.utils import timezone
        self.a2.connections, self.a2.last_seen = 1, timezone.now()
        self.a2.save()
        self.assertEqual(len(self.client.get(reverse("api_employees"), {"online": "1"}).json()["employees"]), 1)
        self.a2.last_seen = timezone.now() - timezone.timedelta(minutes=10)
        self.a2.save()
        self.assertEqual(len(self.client.get(reverse("api_employees"), {"online": "1"}).json()["employees"]), 0)



class AdminTests(TestCase):
    def setUp(self):
        cache.clear()
        self.c, self.other = make_company("A"), make_company("B")
        self.admin = make_user(self.c, "admin@a.com", role=User.ADMIN, first="Boss")
        self.emp = make_user(self.c, "emp@a.com", first="Emp")
        self.stranger = make_user(self.other, "s@b.com", role=User.ADMIN, first="Stranger")
        Message.objects.create(company=self.c, sender=self.emp, recipient=self.admin, body="visible to admin")
        Message.objects.create(company=self.other, sender=self.stranger, recipient=self.stranger, body="other co")

    def test_employee_forbidden(self):
        self.client.force_login(self.emp)
        for name in ("admin_panel", "admin_messages"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 403)
        self.assertEqual(self.client.post(reverse("admin_user_action", args=[self.admin.id]),
                                          {"action": "toggle_active"}).status_code, 403)

    def test_admin_sees_company_data_only(self):
        self.client.force_login(self.admin)
        r = self.client.get(reverse("admin_panel"))
        self.assertContains(r, "emp@a.com")
        self.assertNotContains(r, "s@b.com")
        r = self.client.get(reverse("admin_messages"))
        self.assertContains(r, "visible to admin")
        self.assertNotContains(r, "other co")

    def test_deactivate_and_reactivate_employee(self):
        self.client.force_login(self.admin)
        url = reverse("admin_user_action", args=[self.emp.id])
        self.client.post(url, {"action": "toggle_active"})
        self.emp.refresh_from_db()
        self.assertFalse(self.emp.is_active)
        self.client.post(url, {"action": "toggle_active"})
        self.emp.refresh_from_db()
        self.assertTrue(self.emp.is_active)

    def test_cannot_touch_other_company_user(self):
        self.client.force_login(self.admin)
        r = self.client.post(reverse("admin_user_action", args=[self.stranger.id]), {"action": "toggle_active"})
        self.assertEqual(r.status_code, 404)
        self.stranger.refresh_from_db()
        self.assertTrue(self.stranger.is_active)

    def test_last_admin_protected_and_self_deactivate_blocked(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("admin_user_action", args=[self.admin.id]), {"action": "toggle_active"})
        self.client.post(reverse("admin_user_action", args=[self.admin.id]), {"action": "toggle_role"})
        self.admin.refresh_from_db()
        self.assertTrue(self.admin.is_active)
        self.assertEqual(self.admin.role, User.ADMIN)

    def test_promote_employee(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("admin_user_action", args=[self.emp.id]), {"action": "toggle_role"})
        self.emp.refresh_from_db()
        self.assertEqual(self.emp.role, User.ADMIN)

    def test_message_filters(self):
        self.client.force_login(self.admin)
        self.assertContains(self.client.get(reverse("admin_messages"), {"a": self.emp.id}), "visible to admin")
        self.assertNotContains(self.client.get(reverse("admin_messages"), {"q": "zzz"}), "visible to admin")


class SocketTests(TransactionTestCase):
    """Real-time behaviour through the actual consumer."""

    def setUp(self):
        self.a, self.b = make_company("A"), make_company("B")
        self.u1 = make_user(self.a, "u1@a.com", first="One")
        self.u2 = make_user(self.a, "u2@a.com", first="Two")
        self.x = make_user(self.b, "x@b.com", first="Ex")

    async def _connect(self, user):
        comm = WebsocketCommunicator(URLRouter(websocket_urlpatterns), "/ws/chat/")
        comm.scope["user"] = user
        ok, _ = await comm.connect()
        return comm, ok

    async def _drain(self, comm):
        out = []
        while not await comm.receive_nothing(timeout=0.2):
            out.append(await comm.receive_json_from())
        return out

    async def test_private_message_delivered_and_saved(self):
        c1, ok1 = await self._connect(self.u1)
        c2, ok2 = await self._connect(self.u2)
        self.assertTrue(ok1 and ok2)
        await self._drain(c1)
        await self._drain(c2)
        await c1.send_json_to({"type": "message", "to": self.u2.id, "body": "  hi there  "})
        got = [m for m in await self._drain(c2) if m["type"] == "message"]
        mine = [m for m in await self._drain(c1) if m["type"] == "message"]
        self.assertEqual(got[0]["body"], "hi there")
        self.assertEqual(got[0]["sender"], self.u1.id)
        self.assertEqual(mine[0]["body"], "hi there")  # sender gets an echo too (multi-tab sync)
        self.assertEqual(await Message.objects.acount(), 1)
        await c1.disconnect()
        await c2.disconnect()

    async def test_sender_also_gets_echo(self):
        c1, _ = await self._connect(self.u1)
        await self._drain(c1)
        await c1.send_json_to({"type": "message", "to": self.u2.id, "body": "echo"})
        msgs = [m for m in await self._drain(c1) if m["type"] == "message"]
        self.assertEqual(msgs[0]["body"], "echo")
        await c1.disconnect()

    async def test_cross_company_message_refused(self):
        c1, _ = await self._connect(self.u1)
        cx, _ = await self._connect(self.x)
        await self._drain(c1)
        await self._drain(cx)
        await c1.send_json_to({"type": "message", "to": self.x.id, "body": "leak?"})
        self.assertEqual((await c1.receive_json_from())["type"], "error")
        self.assertEqual([m for m in await self._drain(cx) if m["type"] == "message"], [])
        self.assertEqual(await Message.objects.acount(), 0)
        await c1.disconnect()
        await cx.disconnect()

    async def test_invalid_messages_refused(self):
        c1, _ = await self._connect(self.u1)
        await self._drain(c1)
        for body in ("", "   ", "x" * 2001):
            await c1.send_json_to({"type": "message", "to": self.u2.id, "body": body})
            self.assertEqual((await c1.receive_json_from())["type"], "error")
        await c1.send_json_to({"type": "message", "to": self.u1.id, "body": "to self"})
        self.assertEqual((await c1.receive_json_from())["type"], "error")
        await c1.disconnect()

    async def test_presence_online_then_offline(self):
        c2, _ = await self._connect(self.u2)
        await self._drain(c2)
        c1, _ = await self._connect(self.u1)
        events = await self._drain(c2)
        self.assertIn({"type": "presence", "user_id": self.u1.id, "online": True}, events)
        await c1.disconnect()
        events = await self._drain(c2)
        self.assertIn({"type": "presence", "user_id": self.u1.id, "online": False}, events)
        await c2.disconnect()

    async def test_two_tabs_stay_online_until_last_closes(self):
        c2, _ = await self._connect(self.u2)
        await self._drain(c2)
        t1, _ = await self._connect(self.u1)
        t2, _ = await self._connect(self.u1)
        await self._drain(c2)
        await t1.disconnect()
        self.assertIn({"type": "presence", "user_id": self.u1.id, "online": True}, await self._drain(c2))
        await t2.disconnect()
        self.assertIn({"type": "presence", "user_id": self.u1.id, "online": False}, await self._drain(c2))
        await c2.disconnect()

    async def test_typing_and_cross_company_typing(self):
        c1, _ = await self._connect(self.u1)
        c2, _ = await self._connect(self.u2)
        cx, _ = await self._connect(self.x)
        for c in (c1, c2, cx):
            await self._drain(c)
        await c1.send_json_to({"type": "typing", "to": self.u2.id})
        await c1.send_json_to({"type": "typing", "to": self.x.id})
        self.assertIn({"type": "typing", "from": self.u1.id}, await self._drain(c2))
        self.assertEqual(await self._drain(cx), [])
        for c in (c1, c2, cx):
            await c.disconnect()

    async def test_anonymous_rejected(self):
        from django.contrib.auth.models import AnonymousUser
        comm = WebsocketCommunicator(URLRouter(websocket_urlpatterns), "/ws/chat/")
        comm.scope["user"] = AnonymousUser()
        ok, _ = await comm.connect()
        self.assertFalse(ok)


# ---------------------------------------------------------------- groups
class GroupApiTests(TestCase):
    def setUp(self):
        cache.clear()
        self.a, self.b = make_company("A"), make_company("B")
        self.u1 = make_user(self.a, "u1@a.com", first="One")
        self.u2 = make_user(self.a, "u2@a.com", first="Two")
        self.u3 = make_user(self.a, "u3@a.com", first="Three")
        self.x = make_user(self.b, "x@b.com", first="Ex")
        self.client.force_login(self.u1)

    def post(self, url, data=None):
        return self.client.post(url, json.dumps(data or {}), content_type="application/json")

    def make(self, members=None, name="Team"):
        r = self.post(reverse("api_groups"), {"name": name, "member_ids": members or [self.u2.id, self.u3.id]})
        self.assertEqual(r.status_code, 201, r.content)
        return r.json()["id"]

    def test_create_and_list(self):
        gid = self.make()
        g = Group.objects.get(pk=gid)
        self.assertEqual((g.created_by_id, g.company_id, g.memberships.count()), (self.u1.id, self.a.id, 3))
        lst = self.client.get(reverse("api_groups")).json()["groups"]
        self.assertEqual([(x["name"], x["members"], x["is_owner"]) for x in lst], [("Team", 3, True)])

    def test_create_validation(self):
        url = reverse("api_groups")
        self.assertEqual(self.post(url, {"name": "  ", "member_ids": [self.u2.id]}).status_code, 400)
        self.assertEqual(self.post(url, {"name": "Solo", "member_ids": []}).status_code, 400)
        self.assertEqual(self.post(url, {"name": "Leak", "member_ids": [self.x.id]}).status_code, 400)  # other company
        self.assertEqual(Group.objects.count(), 0)
        gid = self.make([self.u2.id, self.x.id])
        self.assertEqual(GroupMember.objects.filter(group_id=gid).count(), 2)   # x silently not added

    def test_only_members_can_open(self):
        gid = self.make([self.u2.id])
        self.client.force_login(self.u3)
        self.assertEqual(self.client.get(reverse("api_group_detail", args=[gid])).status_code, 404)
        self.client.force_login(self.x)
        self.assertEqual(self.client.get(reverse("api_group_detail", args=[gid])).status_code, 404)
        self.client.force_login(self.u2)
        j = self.client.get(reverse("api_group_detail", args=[gid])).json()
        self.assertEqual(len(j["group"]["members"]), 2)
        self.assertFalse(j["group"]["is_owner"])

    def test_group_messages_in_detail_and_unread(self):
        gid = self.make([self.u2.id])
        Message.objects.create(company=self.a, sender=self.u1, group_id=gid, body="hi team")
        self.client.force_login(self.u2)
        self.assertEqual(self.client.get(reverse("api_groups")).json()["groups"][0]["unread"], 1)
        j = self.client.get(reverse("api_group_detail", args=[gid])).json()
        self.assertEqual(j["messages"][0]["body"], "hi team")
        self.assertEqual(j["messages"][0]["sender_name"], "One User")
        self.assertEqual(self.client.get(reverse("api_groups")).json()["groups"][0]["unread"], 0)

    def test_group_messages_stay_out_of_private_views(self):
        gid = self.make([self.u2.id])
        Message.objects.create(company=self.a, sender=self.u1, group_id=gid, body="group only")
        emp = self.client.get(reverse("api_employees")).json()["employees"]
        self.assertTrue(all(e["last_message"] == "" for e in emp))
        self.assertEqual(self.client.get(reverse("api_messages", args=[self.u2.id])).json()["messages"], [])
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 200)

    def test_add_members_owner_only(self):
        gid = self.make([self.u2.id])
        url = reverse("api_group_add", args=[gid])
        self.assertEqual(self.post(url, {"user_ids": [self.u3.id, self.x.id]}).json()["added"], 1)
        self.client.force_login(self.u2)
        self.assertEqual(self.post(url, {"user_ids": [self.u3.id]}).status_code, 403)

    def test_remove_leave_transfer_and_delete(self):
        gid = self.make([self.u2.id, self.u3.id])
        rm = lambda uid: reverse("api_group_remove", args=[gid, uid])
        self.client.force_login(self.u2)                                   # member cannot remove others
        self.assertEqual(self.post(rm(self.u3.id)).status_code, 403)
        self.assertEqual(self.post(rm(self.u2.id)).status_code, 200)       # but can leave
        self.client.force_login(self.u1)
        self.assertEqual(self.post(rm(self.u3.id)).status_code, 200)       # owner removes
        self.assertEqual(GroupMember.objects.filter(group_id=gid).count(), 1)
        gid2 = self.make([self.u2.id], name="Two")
        self.post(reverse("api_group_remove", args=[gid2, self.u1.id]))     # owner leaves -> transfer
        self.assertEqual(Group.objects.get(pk=gid2).created_by_id, self.u2.id)
        self.client.force_login(self.u2)
        self.post(reverse("api_group_remove", args=[gid2, self.u2.id]))     # last one leaves -> group gone
        self.assertFalse(Group.objects.filter(pk=gid2).exists())

    def test_seen_via_api_marks_read(self):
        Message.objects.create(company=self.a, sender=self.u1, recipient=self.u2, body="yo")
        self.client.force_login(self.u2)
        self.client.get(reverse("api_messages", args=[self.u1.id]))
        self.assertIsNotNone(Message.objects.get(body="yo").read_at)
        self.client.force_login(self.u1)
        m = self.client.get(reverse("api_messages", args=[self.u2.id])).json()["messages"][0]
        self.assertTrue(m["read"])

    def test_admin_audit_shows_group_and_deleted(self):
        self.u1.role = User.ADMIN
        self.u1.save()
        gid = self.make([self.u2.id])
        Message.objects.create(company=self.a, sender=self.u2, group_id=gid, body="group text")
        Message.objects.create(company=self.a, sender=self.u2, group_id=gid, body="", deleted_at=timezone.now())
        r = self.client.get(reverse("admin_messages"))
        self.assertContains(r, "Group: Team")
        self.assertContains(r, "group text")
        self.assertContains(r, "deleted by sender")


# ---------------------------------------------------------------- Excel import + forced password change
def make_xlsx(rows, headers=("Name", "Email", "Job Title", "Department")):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    if headers:
        ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "employees.xlsx"
    return buf


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.PBKDF2PasswordHasher",
                                     "django.contrib.auth.hashers.MD5PasswordHasher"])   # the real hasher
class ImportTests(TestCase):
    def setUp(self):
        cache.clear()
        self.c = make_company("Acme")
        self.admin = make_user(self.c, "admin@acme.com", role=User.ADMIN, first="Boss")
        self.emp = make_user(self.c, "emp@acme.com")
        self.client.force_login(self.admin)
        self.url = reverse("admin_import")

    def upload(self, f):
        return self.client.post(self.url, {"file": f})

    def test_import_creates_accounts_with_temp_passwords(self):
        r = self.upload(make_xlsx([("Priya Sharma", "Priya@Acme.com", "Detailer", "Detailing"),
                                   ("Ravi K", "ravi@acme.com", "", "")]))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.context["created"]), 2)
        self.assertContains(r, "data:application/vnd.openxmlformats")
        u = User.objects.get(username="priya@acme.com")
        self.assertEqual((u.company_id, u.role, u.job_title, u.must_change_password),
                         (self.c.id, User.EMPLOYEE, "Detailer", True))
        pw = r.context["created"][0]["password"]
        self.assertTrue(u.check_password(pw))
        self.assertNotIn(pw, u.password)            # stored hashed, not plain

    def test_skips_bad_duplicate_and_existing_rows(self):
        r = self.upload(make_xlsx([("Good One", "good@acme.com", "", ""), ("Dup", "GOOD@acme.com", "", ""),
                                   ("Bad Mail", "not-an-email", "", ""), ("", "noname@acme.com", "", ""),
                                   ("Existing", "emp@acme.com", "", "")]))
        self.assertEqual([c["email"] for c in r.context["created"]], ["good@acme.com"])
        reasons = sorted(s["reason"] for s in r.context["skipped"])
        self.assertEqual(len(reasons), 4)
        self.assertTrue(any("already has an account" in x for x in reasons))
        self.assertTrue(any("Duplicate" in x for x in reasons))

    def test_headings_are_flexible_and_required(self):
        r = self.upload(make_xlsx([("Meena", "meena@acme.com", "HR")], headers=("Employee Name", "E-mail", "Designation")))
        self.assertEqual(User.objects.get(username="meena@acme.com").job_title, "HR")
        r = self.upload(make_xlsx([("x", "y")], headers=("Foo", "Bar")))
        self.assertContains(r, "must have the headings")

    def test_emp_code_becomes_login_id_and_headings_are_loose(self):
        f = make_xlsx([("EPRO0023", "Veeresh Hottimath", "pradeep@acme.com", ""),
                       ("EPRO0024", "Ameet Paradeshi", "", ""),
                       ("EPRO0039", "Venkanna Madivalar", "venkannamadivalar_gmail.com#EXT#@acme.onmicrosoft.com", "")],
                      headers=("Emp Code", "NAME", "Email.id"))
        r = self.upload(f)
        self.assertEqual([c["login"] for c in r.context["created"]], ["epro0023", "epro0024", "epro0039"])
        u = User.objects.get(username="epro0039")
        self.assertEqual(u.email, "venkannamadivalar@gmail.com")             # Microsoft guest id turned into the real email
        self.assertEqual(User.objects.get(username="epro0024").email, "")    # email is optional
        self.assertEqual(User.objects.get(username="epro0023").email, "")    # pradeep@ is not Veeresh Hottimath's email: ignored
        self.assertEqual([bool(c["warn"]) for c in r.context["created"]], [True, False, False])
        self.assertEqual(r.context["email_ignored"], 1)
        self.assertContains(r, "were ignored because they do not look like the person")
        cache.clear()
        self.client.logout()
        r = self.client.post(reverse("login"), {"username": "EPRO0024", "password": r.context["created"][1]["password"]})
        self.assertRedirects(r, reverse("dashboard"), fetch_redirect_response=False)

    def test_repeated_or_missing_emp_codes_are_skipped_not_guessed(self):
        f = make_xlsx([("EPRO", "Srivatsan", "", ""), ("EPRO", "Harish", "", ""), ("EPRO0100", "Santhosh M", "", ""),
                       ("", "No Code Person", "", ""), ("", "", "orphan@acme.com", "")],
                      headers=("Emp Code", "Name", "Email"))
        r = self.upload(f)
        self.assertEqual([c["login"] for c in r.context["created"]], ["epro0100"])
        reasons = {s["name"]: s["reason"] for s in r.context["skipped"]}
        self.assertIn("more than one row", reasons["Srivatsan"])
        self.assertIn("more than one row", reasons["Harish"])
        self.assertEqual(reasons["No Code Person"], "No Emp Code or Email")
        self.assertEqual(reasons[""], "Name is missing or too long")

    def test_rejects_wrong_file_and_too_many_rows(self):
        self.assertContains(self.client.post(self.url, {"file": io.BytesIO(b"hello")}), "ending in .xlsx")
        bad = io.BytesIO(b"not a real workbook")
        bad.name = "x.xlsx"
        self.assertContains(self.upload(bad), "Could not read this file")
        many = make_xlsx([(f"P{i}", f"p{i}@acme.com", "", "") for i in range(301)])
        self.assertContains(self.upload(many), "Too many rows")
        self.assertEqual(User.objects.filter(company=self.c).count(), 2)

    def test_employee_cannot_import(self):
        self.client.force_login(self.emp)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.upload(make_xlsx([("A", "a@acme.com", "", "")])).status_code, 403)
        self.assertEqual(self.client.get(reverse("admin_import_template")).status_code, 403)

    def test_import_only_into_own_company(self):
        other = make_company("Other")
        self.upload(make_xlsx([("New", "new@acme.com", "", "")]))
        self.assertEqual(User.objects.get(username="new@acme.com").company_id, self.c.id)
        self.assertEqual(other.users.count(), 0)

    def test_template_download_is_a_valid_workbook(self):
        from openpyxl import load_workbook
        r = self.client.get(reverse("admin_import_template"))
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual([c.value for c in wb["Employees"][1]], ["Emp Code", "Name", "Email", "Job Title", "Department"])
        self.assertEqual(wb["Employees"].max_row, 1)       # template itself has no sample employee to import by accident

    def test_credentials_export_never_makes_formulas(self):
        from openpyxl import load_workbook
        from . import importer
        data = importer.credentials_xlsx([{"name": "=1+1", "login": "e1", "email": "e@acme.com", "password": "pw",
                                           "job_title": "@x", "department": "+y"}])
        ws = load_workbook(io.BytesIO(data)).active
        self.assertEqual(ws["A2"].data_type, "s")
        self.assertEqual(ws["A2"].value, "=1+1")

    def test_imported_user_must_change_password_then_can_chat(self):
        r = self.upload(make_xlsx([("New Hire", "hire@acme.com", "", "")]))
        pw = r.context["created"][0]["password"]
        self.client.logout()
        self.assertRedirects(self.client.post(reverse("login"), {"username": "hire@acme.com", "password": pw}),
                             reverse("dashboard"), fetch_redirect_response=False)
        for name in ("dashboard", "chat", "directory"):
            self.assertRedirects(self.client.get(reverse(name)), reverse("password_change"))
        self.assertEqual(self.client.get(reverse("api_employees")).status_code, 403)
        self.assertEqual(self.client.get(reverse("password_change")).status_code, 200)
        bad = self.client.post(reverse("password_change"), {"old_password": pw, "new_password1": "12345678",
                                                            "new_password2": "12345678"})
        self.assertEqual(bad.status_code, 200)
        self.assertTrue(User.objects.get(username="hire@acme.com").must_change_password)
        ok = self.client.post(reverse("password_change"), {"old_password": pw, "new_password1": "My-New-Pass-77",
                                                           "new_password2": "My-New-Pass-77"})
        self.assertRedirects(ok, reverse("dashboard"))
        self.assertFalse(User.objects.get(username="hire@acme.com").must_change_password)
        self.assertEqual(self.client.get(reverse("chat")).status_code, 200)

    def test_temp_password_hash_is_light_then_upgraded_at_first_login(self):
        r = self.upload(make_xlsx([("Hash Test", "hash@acme.com", "", "")]))
        pw = r.context["created"][0]["password"]
        u = User.objects.get(username="hash@acme.com")
        self.assertTrue(u.password.startswith("pbkdf2_sha256$60000$"))
        self.client.logout()
        self.client.post(reverse("login"), {"username": "hash@acme.com", "password": pw})
        u.refresh_from_db()
        self.assertFalse(u.password.startswith("pbkdf2_sha256$60000$"))      # re-hashed at full strength
        self.assertTrue(u.check_password(pw))

    def test_normal_user_can_change_password(self):
        self.client.force_login(self.emp)
        r = self.client.post(reverse("password_change"), {"old_password": "Str0ng-pass-99",
                                                          "new_password1": "Another-Pass-55", "new_password2": "Another-Pass-55"})
        self.assertRedirects(r, reverse("dashboard"))


# ---------------------------------------------------------------- live group / seen / delete
class LiveFeatureTests(TransactionTestCase):
    def setUp(self):
        self.a, self.b = make_company("A"), make_company("B")
        self.u1 = make_user(self.a, "u1@a.com", first="One")
        self.u2 = make_user(self.a, "u2@a.com", first="Two")
        self.u3 = make_user(self.a, "u3@a.com", first="Three")
        self.x = make_user(self.b, "x@b.com", first="Ex")
        self.g = Group.objects.create(company=self.a, name="Team", created_by=self.u1)
        for u in (self.u1, self.u2):
            GroupMember.objects.create(group=self.g, user=u)

    async def _connect(self, user):
        comm = WebsocketCommunicator(URLRouter(websocket_urlpatterns), "/ws/chat/")
        comm.scope["user"] = user
        ok, _ = await comm.connect()
        assert ok
        return comm

    async def _drain(self, comm):
        out = []
        while not await comm.receive_nothing(timeout=0.2):
            out.append(await comm.receive_json_from())
        return out

    async def test_group_message_reaches_members_only(self):
        c1, c2, c3 = await self._connect(self.u1), await self._connect(self.u2), await self._connect(self.u3)
        for c in (c1, c2, c3):
            await self._drain(c)
        await c1.send_json_to({"type": "group_message", "group": self.g.id, "body": "hello team"})
        got2 = [m for m in await self._drain(c2) if m["type"] == "message"]
        got1 = [m for m in await self._drain(c1) if m["type"] == "message"]
        self.assertEqual((got2[0]["body"], got2[0]["group"], got2[0]["sender_name"], got2[0]["group_name"]),
                         ("hello team", self.g.id, "One User", "Team"))
        self.assertEqual(got1[0]["body"], "hello team")
        self.assertEqual([m for m in await self._drain(c3) if m["type"] == "message"], [])
        m = await Message.objects.aget(body="hello team")
        self.assertEqual((m.group_id, m.recipient_id), (self.g.id, None))
        for c in (c1, c2, c3):
            await c.disconnect()

    async def test_non_member_and_other_company_cannot_post_to_group(self):
        c3, cx = await self._connect(self.u3), await self._connect(self.x)
        for c in (c3, cx):
            await self._drain(c)
            await c.send_json_to({"type": "group_message", "group": self.g.id, "body": "sneaky"})
            self.assertEqual((await c.receive_json_from())["type"], "error")
        self.assertEqual(await Message.objects.acount(), 0)
        for c in (c3, cx):
            await c.disconnect()

    async def test_delete_own_message_only(self):
        c1, c2 = await self._connect(self.u1), await self._connect(self.u2)
        for c in (c1, c2):
            await self._drain(c)
        await c1.send_json_to({"type": "message", "to": self.u2.id, "body": "oops secret"})
        await self._drain(c1)
        await self._drain(c2)
        mid = (await Message.objects.aget(body="oops secret")).id
        await c2.send_json_to({"type": "delete", "id": mid})                     # receiver cannot delete it
        self.assertEqual((await c2.receive_json_from())["type"], "error")
        self.assertTrue(await Message.objects.filter(pk=mid, deleted_at__isnull=True).aexists())
        await c1.send_json_to({"type": "delete", "id": mid})                     # sender can
        ev1 = [m for m in await self._drain(c1) if m["type"] == "deleted"]
        ev2 = [m for m in await self._drain(c2) if m["type"] == "deleted"]
        self.assertEqual((ev1[0]["id"], ev2[0]["id"]), (mid, mid))
        m = await Message.objects.aget(pk=mid)
        self.assertEqual((m.body, m.deleted_at is not None), ("", True))         # text really erased
        await c1.send_json_to({"type": "delete", "id": mid})                     # already deleted
        self.assertEqual((await c1.receive_json_from())["type"], "error")
        for c in (c1, c2):
            await c.disconnect()

    async def test_delete_group_message_reaches_all_members(self):
        c1, c2 = await self._connect(self.u1), await self._connect(self.u2)
        await c2.send_json_to({"type": "group_message", "group": self.g.id, "body": "bye"})
        await self._drain(c1)
        await self._drain(c2)
        mid = (await Message.objects.aget(body="bye")).id
        await c2.send_json_to({"type": "delete", "id": mid})
        self.assertIn(mid, [m["id"] for m in await self._drain(c1) if m["type"] == "deleted"])
        await c1.send_json_to({"type": "delete", "id": mid})                     # not the sender
        self.assertEqual((await c1.receive_json_from())["type"], "error")
        for c in (c1, c2):
            await c.disconnect()

    async def test_seen_event_when_recipient_reads(self):
        c1, c2 = await self._connect(self.u1), await self._connect(self.u2)
        for c in (c1, c2):
            await self._drain(c)
        await c1.send_json_to({"type": "message", "to": self.u2.id, "body": "read me"})
        await self._drain(c1)
        await self._drain(c2)
        await c2.send_json_to({"type": "read", "peer": self.u1.id})
        self.assertIn({"type": "seen", "by": self.u2.id}, await self._drain(c1))
        self.assertIsNotNone((await Message.objects.aget(body="read me")).read_at)
        await c2.send_json_to({"type": "read", "peer": self.u1.id})              # nothing new -> no second event
        self.assertEqual([m for m in await self._drain(c1) if m["type"] == "seen"], [])
        for c in (c1, c2):
            await c.disconnect()

    async def test_seen_event_when_recipient_opens_chat_over_http(self):
        from asgiref.sync import sync_to_async
        from django.test import Client
        c1 = await self._connect(self.u1)
        await self._drain(c1)
        await Message.objects.acreate(company=self.a, sender=self.u1, recipient=self.u2, body="hi")
        def open_chat():
            cl = Client()
            cl.force_login(self.u2)
            return cl.get(reverse("api_messages", args=[self.u1.id])).status_code
        self.assertEqual(await sync_to_async(open_chat)(), 200)
        self.assertIn({"type": "seen", "by": self.u2.id}, await self._drain(c1))
        await c1.disconnect()

    async def test_must_change_password_cannot_connect(self):
        await User.objects.filter(pk=self.u3.pk).aupdate(must_change_password=True)
        u3 = await User.objects.aget(pk=self.u3.pk)
        comm = WebsocketCommunicator(URLRouter(websocket_urlpatterns), "/ws/chat/")
        comm.scope["user"] = u3
        ok, _ = await comm.connect()
        self.assertFalse(ok)


# ---------------------------------------------------------------- no self-registration + default admin
class NoRegistrationTests(TestCase):
    def test_register_page_is_gone(self):
        r = self.client.get("/register/")
        self.assertRedirects(r, reverse("login"), fetch_redirect_response=False)
        r = self.client.post("/register/", {"mode": "create", "company_name": "Evil", "full_name": "Eve Il",
                                            "email": "eve@evil.com", "password1": "Str0ng-pass-99",
                                            "password2": "Str0ng-pass-99"})
        self.assertEqual(r.status_code, 302)
        self.assertFalse(User.objects.exists() or Company.objects.exists())     # nothing was created

    def test_login_page_has_no_register_links(self):
        r = self.client.get(reverse("login"))
        self.assertNotContains(r, "Create a company")
        self.assertNotContains(r, "join with a code")
        self.assertContains(r, "Login ID")

    def test_login_page_has_show_hide_password_button(self):
        r = self.client.get(reverse("login"))
        self.assertContains(r, 'id="pwtoggle"')
        self.assertContains(r, 'type="password"')

    def test_can_log_in_with_a_login_id_that_is_not_an_email(self):
        cache.clear()
        u = make_user(make_company("Acme"), "bharath", role=User.ADMIN, first="Bharath")
        r = self.client.post(reverse("login"), {"username": "  BHARATH ", "password": "Str0ng-pass-99"})
        self.assertRedirects(r, reverse("dashboard"), fetch_redirect_response=False)
        self.assertEqual(int(self.client.session["_auth_user_id"]), u.id)


class DefaultAdminTests(TestCase):
    def run_it(self, **env):
        from unittest import mock
        from . import bootstrap
        lines = []
        with mock.patch.dict("os.environ", env, clear=False):
            for k in ("DEFAULT_ADMIN_USERNAME", "DEFAULT_ADMIN_NAME", "DEFAULT_ADMIN_PASSWORD", "DEFAULT_COMPANY_NAME",
                      "DEFAULT_ADMIN_RESET", "DISABLE_DEFAULT_ADMIN"):
                if k not in env:
                    import os
                    os.environ.pop(k, None)
            user = bootstrap.ensure_default_admin(out=lines.append)
        return user, lines

    def test_creates_admin_with_the_owners_own_password(self):
        u, lines = self.run_it(DEFAULT_ADMIN_PASSWORD="bharath@2916")
        self.assertEqual((u.username, u.first_name, u.role, u.must_change_password), ("bharath", "Bharath", User.ADMIN, False))
        self.assertTrue(u.check_password("bharath@2916"))
        self.assertEqual(u.company.name, "Bharath")
        self.assertNotIn("bharath@2916", " ".join(lines))            # a password from env is never printed
        cache.clear()
        r = self.client.post(reverse("login"), {"username": "Bharath", "password": "bharath@2916"})
        self.assertRedirects(r, reverse("dashboard"), fetch_redirect_response=False)
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 200)   # no forced change for the owner's choice

    def test_random_password_is_printed_once_when_none_given(self):
        u, lines = self.run_it()
        pw = lines[0].split("Temporary password: ")[1].split(" ")[0]
        self.assertEqual(len(pw), 14)
        self.assertTrue(u.check_password(pw))
        self.assertTrue(u.must_change_password)                      # a random password must be replaced

    def test_is_idempotent_and_never_touches_an_existing_account(self):
        u, _ = self.run_it(DEFAULT_ADMIN_PASSWORD="First-Pass-123")
        u.set_password("My-Own-Pass-77")
        u.must_change_password = False
        u.last_login = timezone.now()                                # the owner has logged in at least once
        u.save()
        again, lines = self.run_it(DEFAULT_ADMIN_PASSWORD="Other-Pass-999")
        self.assertIsNone(again)
        self.assertEqual(lines, [])
        u.refresh_from_db()
        self.assertTrue(u.check_password("My-Own-Pass-77"))
        self.assertEqual(User.objects.count(), 1)

    def test_changed_password_in_render_applies_until_first_login(self):
        u, _ = self.run_it()                                         # first deploy without a password: random one
        self.assertTrue(u.must_change_password)
        _, lines = self.run_it(DEFAULT_ADMIN_PASSWORD="bharath@2916")  # owner then sets the password in Render
        u.refresh_from_db()
        self.assertTrue(u.check_password("bharath@2916"))
        self.assertFalse(u.must_change_password)
        self.assertEqual(len(lines), 1)
        _, lines = self.run_it(DEFAULT_ADMIN_PASSWORD="bharath@2916")  # same password again: nothing to do
        self.assertEqual(lines, [])

    def test_reset_flag_recovers_a_forgotten_password(self):
        u, _ = self.run_it(DEFAULT_ADMIN_PASSWORD="First-Pass-123")
        u.is_active = False
        u.must_change_password = False
        u.save()
        self.run_it(DEFAULT_ADMIN_PASSWORD="Recover-Pass-55", DEFAULT_ADMIN_RESET="1")
        u.refresh_from_db()
        self.assertTrue(u.check_password("Recover-Pass-55"))
        self.assertTrue(u.is_active)
        self.assertFalse(u.must_change_password)                      # password given by the owner

    def test_joins_the_only_existing_company(self):
        existing = make_company("Bharath Steel")
        u, _ = self.run_it(DEFAULT_ADMIN_PASSWORD="Given-Pass-123")
        self.assertEqual(u.company_id, existing.id)
        self.assertEqual(Company.objects.count(), 1)

    def test_custom_names_and_disable_switch(self):
        u, _ = self.run_it(DEFAULT_ADMIN_USERNAME="Boss", DEFAULT_ADMIN_NAME="Big Boss",
                           DEFAULT_COMPANY_NAME="Edge Pro", DEFAULT_ADMIN_PASSWORD="Given-Pass-123")
        self.assertEqual((u.username, u.first_name, u.last_name, u.company.name), ("boss", "Big", "Boss", "Edge Pro"))
        User.objects.all().delete()
        none, _ = self.run_it(DISABLE_DEFAULT_ADMIN="1")
        self.assertIsNone(none)
        self.assertFalse(User.objects.exists())
