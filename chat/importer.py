"""Excel import of employees for a company admin.

Columns (headings are matched loosely: capitals, spaces and dots do not matter):
  Name        required
  Emp Code    the Login ID (e.g. EPRO0023)   \\  one of these two is required;
  Email       optional contact email          /  the Login ID is the Emp Code, else the email
  Job Title, Department   optional
"""
import io
import re
import secrets

from django.contrib.auth.hashers import PBKDF2PasswordHasher
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from openpyxl import Workbook, load_workbook

from .models import User

MAX_ROWS = 300
MAX_BYTES = 2 * 1024 * 1024
HEADERS = {
    "name": {"name", "fullname", "employeename", "employee", "staffname"},
    "login": {"empcode", "employeecode", "empid", "employeeid", "staffid", "loginid", "userid", "username", "code"},
    "email": {"email", "emailid", "workemail", "mail", "emailaddress"},
    "job_title": {"jobtitle", "title", "designation", "role", "position"},
    "department": {"department", "dept", "team"},
}
_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghjkmnpqrstuvwxyz23456789"
_LOGIN_OK = re.compile(r"^[a-z0-9][a-z0-9._@+-]{1,149}$")
_GUEST = re.compile(r"^(?P<local>.+)_(?P<domain>[^_@#]+\.[a-z0-9.]+)#EXT#@", re.I)   # Microsoft 365 guest id


class ImportProblem(Exception):
    """A problem with the whole file (shown to the admin as one message)."""


class _QuickHasher(PBKDF2PasswordHasher):
    # Temporary passwords are random and must be replaced at first login. A lighter hash keeps a
    # 300-person import fast; Django re-hashes with the full strength the first time the person logs in.
    iterations = 60000


def _hash_temp(password):
    h = _QuickHasher()
    return h.encode(password, h.salt())


def temp_password():
    return "".join(secrets.choice(_ALPHABET) for _ in range(10))


def _norm(heading):
    return re.sub(r"[^a-z0-9]", "", str(heading or "").lower())


def read_rows(fileobj):
    try:
        wb = load_workbook(fileobj, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            raise ImportProblem("The first sheet is empty.")
        cols = {}
        for idx, h in enumerate(header):
            key = _norm(h)
            for field, aliases in HEADERS.items():
                if key in aliases and field not in cols:
                    cols[field] = idx
        if "name" not in cols or ("login" not in cols and "email" not in cols):
            raise ImportProblem("Row 1 must have the headings Name and either Emp Code (the Login ID) or Email "
                                "(optional: Job Title, Department). Download the template to see the format.")
        rows = []
        for n, r in enumerate(rows_iter, start=2):
            vals = {f: (str(r[i]).strip() if i < len(r) and r[i] is not None else "") for f, i in cols.items()}
            if not any(vals.values()):
                continue
            vals["row"] = n
            rows.append(vals)
            if len(rows) > MAX_ROWS:
                raise ImportProblem(f"Too many rows. Import at most {MAX_ROWS} employees per file; split the file.")
        if not rows:
            raise ImportProblem("No employee rows found under the headings.")
        return rows
    except ImportProblem:
        raise
    except Exception:
        raise ImportProblem("Could not read this file. Upload an Excel .xlsx file.") from None


def clean_email(raw):
    """A usable contact email, or ''. Turns Microsoft guest ids (x_gmail.com#EXT#@tenant...) into x@gmail.com."""
    e = (raw or "").strip().lower().replace(" ", "")
    m = _GUEST.match(e)
    if m:
        e = f"{m.group('local')}@{m.group('domain')}"
    try:
        validate_email(e)
    except ValidationError:
        return ""
    return e if len(e) <= 150 else ""


def _looks_like(name, email):
    """Rough check that an email belongs to this name (first 4 letters of any name part appear in it)."""
    local = email.split("@")[0]
    parts = [p for p in re.split(r"[^a-z]+", name.lower()) if len(p) > 2]
    return (not parts) or any(p[:4] in local for p in parts)


def import_employees(company, rows):
    """Create accounts. Returns (created, skipped); created rows carry the temporary password."""
    created, skipped, seen = [], [], set()
    code_counts = {}
    for r in rows:
        c = (r.get("login") or "").strip().lower()
        if c:
            code_counts[c] = code_counts.get(c, 0) + 1
    for r in rows:
        name = " ".join(r.get("name", "").split())
        code = (r.get("login") or "").strip().lower()
        email = clean_email(r.get("email", ""))
        who = {"row": r["row"], "name": name, "login": code, "email": email}
        if not name or len(name) > 150:
            skipped.append({**who, "reason": "Name is missing or too long"}); continue
        if code:
            if code_counts[code] > 1:
                skipped.append({**who, "reason": f"Emp Code '{code.upper()}' is used on more than one row"}); continue
            login = code
        elif email:
            login = email
        else:
            skipped.append({**who, "reason": ("Email is not valid" if r.get("email") else "No Emp Code or Email")}); continue
        who["login"] = login
        if not _LOGIN_OK.match(login):
            skipped.append({**who, "reason": "Login ID has invalid characters"}); continue
        if login in seen:
            skipped.append({**who, "reason": "Duplicate Login ID in this file"}); continue
        seen.add(login)
        if User.objects.filter(username__iexact=login).exists():
            skipped.append({**who, "reason": "This Login ID already has an account"}); continue
        warn = ""
        if email and not _looks_like(name, email):
            email, warn = "", "Email ignored - it does not look like this person's name"
            who["email"] = ""
        first, _, last = name.partition(" ")
        pw = temp_password()
        user = User(username=login, email=email, first_name=first[:150], last_name=last[:150], company=company,
                    role=User.EMPLOYEE, job_title=r.get("job_title", "")[:100],
                    department=r.get("department", "")[:100], must_change_password=True)
        user.password = _hash_temp(pw)
        user.save()
        created.append({**who, "job_title": user.job_title, "department": user.department,
                        "password": pw, "warn": warn})
    return created, skipped


def _put(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"   # text only - never let Excel treat imported text as a formula
    return cell


def credentials_xlsx(created):
    wb = Workbook()
    ws = wb.active
    ws.title = "New accounts"
    for c, h in enumerate(["Name", "Login ID", "Temporary password", "Email", "Job title", "Department"], 1):
        ws.cell(row=1, column=c, value=h)
    for r, u in enumerate(created, 2):
        for c, k in enumerate(["name", "login", "password", "email", "job_title", "department"], 1):
            _put(ws, r, c, u[k])
    for col, w in zip("ABCDEF", (28, 18, 20, 34, 24, 20)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    for c, h in enumerate(["Emp Code", "Name", "Email", "Job Title", "Department"], 1):
        ws.cell(row=1, column=c, value=h)
    for col, w in zip("ABCDE", (14, 28, 34, 24, 20)):
        ws.column_dimensions[col].width = w
    ex = wb.create_sheet("Example")
    ex.append(["Emp Code", "Name", "Email", "Job Title", "Department"])
    ex.append(["EPRO0023", "Priya Sharma", "priya@yourcompany.com", "Senior Detailer", "Detailing"])
    ex.append(["Emp Code becomes the Login ID. Email is optional. Fill in the 'Employees' sheet only."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
